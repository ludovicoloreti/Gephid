"""Smoke test dei formati documentali sulle librerie realmente installate."""
import io
from email.message import EmailMessage

import docx
import openpyxl
import pymupdf
import pytest

import diffuchat as D


@pytest.fixture(autouse=True)
def isolated_ingest(monkeypatch):
    monkeypatch.setattr(D, "INGEST", {})


def test_word_preserves_paragraphs_and_tables():
    document = docx.Document()
    document.add_paragraph("Contratto ZA FFIRO")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Penale"
    table.cell(0, 1).text = "15%"
    data = io.BytesIO()
    document.save(data)
    result = D.ingest_file("contratto.docx", data.getvalue())
    text = D.INGEST[result["id"]]["text"]
    assert result["kind"] == "doc" and "Contratto ZA FFIRO" in text and "Penale\t15%" in text


def test_excel_reads_sheet_names_and_cell_values():
    workbook = openpyxl.Workbook()
    workbook.active.title = "Fatture"
    workbook.active.append(["Cliente", "Totale"])
    workbook.active.append(["Zaffiro", 123.45])
    data = io.BytesIO()
    workbook.save(data)
    result = D.ingest_file("fatture.xlsx", data.getvalue())
    text = D.INGEST[result["id"]]["text"]
    assert "Fatture" in text and "Zaffiro\t123.45" in text


def test_email_extracts_html_body_without_scripts():
    message = EmailMessage()
    message["Subject"] = "Ordine Zaffiro"
    message.set_content("<p>Totale: <b>120 euro</b></p><script>secret()</script>", subtype="html")
    result = D.ingest_file("ordine.eml", message.as_bytes())
    text = D.INGEST[result["id"]]["text"]
    assert "Ordine Zaffiro" in text and "120 euro" in text and "secret" not in text


def test_pdf_export_is_readable_and_can_be_attached():
    data = D._html_to_pdf("<h1>Documento Zaffiro</h1><p>Il codice del progetto e ZA FFIRO-7731.</p>", "Gephid export")
    with pymupdf.open(stream=data, filetype="pdf") as document:
        assert len(document) == 1
        assert "Gephid export" in document[0].get_text()
    result = D.ingest_file("esportazione.pdf", data)
    assert result["kind"] == "doc" and "7731" in D.INGEST[result["id"]]["text"]


@pytest.mark.parametrize("filename", ["note.txt", "note.md", "dati.csv", "script.py"])
def test_text_attachments_preserve_unicode(filename):
    result = D.ingest_file(filename, "Città, quantità: 42\nZAFFIRO-7731".encode())
    assert result["kind"] == "doc" and D.INGEST[result["id"]]["text"] == "Città, quantità: 42\nZAFFIRO-7731"
