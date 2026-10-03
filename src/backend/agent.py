"""Strumenti dell'agente Gephid. Tutti offline: nessuno tocca la rete.

- utility (senza conferma): calcolatrice esatta, data/ora, ricerca nei documenti della chat;
- file, SOLO nella cartella di lavoro autorizzata dall'utente: elenca/leggi/cerca senza conferma,
  scrivi con anteprima + conferma;
- codice (sempre con conferma): Python (o Node se installato) dentro `sandbox-exec`, senza rete in
  uscita (solo localhost), scrittura solo in una cartella temporanea, con timeout; i processi lunghi
  (es. un server) restano in background ma possono ascoltare SOLO su 127.0.0.1.

L'output degli strumenti è NON fidato (può contenere prompt injection da file/documenti): per questo
scrittura ed esecuzione restano sempre dietro un clic dell'utente.
"""
import ast, datetime, decimal, glob, math, os, re, shutil, subprocess, tempfile, threading, time, uuid

# ---------------------------------------------------------------- chiamate native del template
# Il modello emette: call:nome{chiave:<|"|>stringa<|"|>,n:3,l:[...],o:{...}}
_Q = '<|"|>'


def parse_call(raw):
    """'call:nome{...}' -> (nome, dict argomenti). ValueError se non è una chiamata valida."""
    raw = (raw or "").strip()
    m = re.match(r"call:([A-Za-z_][\w]*)\s*\{", raw)
    if not m or not raw.endswith("}"):
        raise ValueError("chiamata a strumento non valida")
    p = _ArgParser(raw, m.end() - 1)
    args = p.obj()
    return m.group(1), args


class _ArgParser:
    def __init__(self, s, i):
        self.s, self.i = s, i

    def ws(self):
        while self.i < len(self.s) and self.s[self.i] in " \n\t\r":
            self.i += 1

    def value(self):
        self.ws()
        s, i = self.s, self.i
        if s.startswith(_Q, i):
            j = s.find(_Q, i + len(_Q))
            if j < 0:
                raise ValueError("stringa non chiusa")
            self.i = j + len(_Q)
            return s[i + len(_Q):j]
        if s[i] == "{":
            return self.obj()
        if s[i] == "[":
            self.i += 1; out = []
            self.ws()
            if s[self.i] == "]":
                self.i += 1; return out
            while True:
                out.append(self.value()); self.ws()
                if s[self.i] == ",": self.i += 1; continue
                if s[self.i] == "]": self.i += 1; return out
                raise ValueError("lista malformata")
        m = re.compile(r"[^,}\]]+").match(s, i)
        if not m:
            raise ValueError("valore mancante")
        self.i = m.end()
        tok = m.group(0).strip()
        if tok in ("true", "false"): return tok == "true"
        if tok in ("null", "none", "None"): return None
        try:
            return int(tok)
        except ValueError:
            try: return float(tok)
            except ValueError: return tok

    def obj(self):
        s = self.s
        assert s[self.i] == "{"
        self.i += 1; out = {}
        self.ws()
        if s[self.i] == "}":
            self.i += 1; return out
        while True:
            self.ws()
            m = re.compile(r"[A-Za-z_][\w]*").match(s, self.i)
            if not m:
                raise ValueError("chiave mancante")
            key = m.group(0); self.i = m.end(); self.ws()
            if s[self.i] != ":":
                raise ValueError("':' mancante")
            self.i += 1
            out[key] = self.value(); self.ws()
            if s[self.i] == ",": self.i += 1; continue
            if s[self.i] == "}": self.i += 1; return out
            raise ValueError("oggetto malformato")


# ---------------------------------------------------------------- calcolatrice esatta (niente eval)
_FUNCS = {"sqrt": lambda x: x.sqrt(), "abs": abs, "round": lambda x, n=0: round(x, int(n)),
          "ln": lambda x: x.ln(), "log10": lambda x: x.log10(), "exp": lambda x: x.exp()}
_FLOAT_FUNCS = {"sin": math.sin, "cos": math.cos, "tan": math.tan, "log": math.log}
_CONST = {"pi": decimal.Decimal("3.14159265358979323846264338327950288"), "e": decimal.Decimal(1).exp()}


def calc(expr):
    """Valuta un'espressione aritmetica con decimali esatti: 0.1+0.2 = 0.3. ValueError se non ammessa."""
    if len(expr) > 300:
        raise ValueError("espressione troppo lunga")
    try:
        tree = ast.parse(expr.replace("^", "**").replace(",", ".") if expr.count(",") and "(" not in expr else expr.replace("^", "**"), mode="eval")
    except SyntaxError:
        raise ValueError("espressione non valida")
    D = decimal.Decimal
    ctx = decimal.Context(prec=28)

    def ev(n):
        if isinstance(n, ast.Expression): return ev(n.body)
        if isinstance(n, ast.Constant) and isinstance(n.value, (int, float)) and not isinstance(n.value, bool):
            return D(str(n.value))
        if isinstance(n, ast.Name) and n.id in _CONST: return _CONST[n.id]
        if isinstance(n, ast.UnaryOp) and isinstance(n.op, (ast.UAdd, ast.USub)):
            v = ev(n.operand); return v if isinstance(n.op, ast.UAdd) else -v
        if isinstance(n, ast.BinOp):
            a, b = ev(n.left), ev(n.right)
            op = n.op
            if isinstance(op, ast.Add): return ctx.add(a, b)
            if isinstance(op, ast.Sub): return ctx.subtract(a, b)
            if isinstance(op, ast.Mult): return ctx.multiply(a, b)
            if isinstance(op, ast.Div): return ctx.divide(a, b)
            if isinstance(op, ast.FloorDiv): return ctx.divide_int(a, b)
            if isinstance(op, ast.Mod): return ctx.remainder(a, b)
            if isinstance(op, ast.Pow):
                if abs(b) > 1000 or abs(a) > D(10) ** 100:
                    raise ValueError("potenza troppo grande")
                return ctx.power(a, b)
        if isinstance(n, ast.Call) and isinstance(n.func, ast.Name) and not n.keywords:
            args = [ev(x) for x in n.args]
            if n.func.id in _FUNCS: return D(_FUNCS[n.func.id](*args))
            if n.func.id in _FLOAT_FUNCS: return D(repr(_FLOAT_FUNCS[n.func.id](*map(float, args))))
        raise ValueError("operazione non ammessa")

    try:
        v = ev(tree)
    except (decimal.InvalidOperation, ZeroDivisionError, OverflowError, TypeError) as e:
        raise ValueError("calcolo impossibile: " + type(e).__name__)
    if v == v.to_integral_value() and abs(v) < D(10) ** 28:
        return str(v.quantize(D(1)))
    s = format(v.normalize(ctx), "f")
    return s


# ---------------------------------------------------------------- ricerca nei documenti della chat
_WORD = re.compile(r"\w{2,}", re.UNICODE)


def _chunks(text, size=900, overlap=150):
    out, i = [], 0
    while i < len(text):
        out.append(text[i:i + size]); i += size - overlap
    return out


def search_docs(docs, query, k=4):
    """BM25 sui pezzi dei documenti allegati alla chat. docs: [{"name","text"}] -> [{"doc","text","score"}]."""
    q = [w.lower() for w in _WORD.findall(query or "")]
    if not q:
        return []
    pieces = [(d["name"], c) for d in docs for c in _chunks(d.get("text") or "")]
    if not pieces:
        return []
    toks = [[w.lower() for w in _WORD.findall(c)] for _, c in pieces]
    N = len(pieces); avg = sum(map(len, toks)) / N or 1
    df = {w: sum(1 for t in toks if w in t) for w in set(q)}
    res = []
    for (name, text), t in zip(pieces, toks):
        tf = {}
        for w in t:
            if w in df: tf[w] = tf.get(w, 0) + 1
        sc = sum(math.log(1 + (N - df[w] + .5) / (df[w] + .5)) * tf[w] * 2.2 / (tf[w] + 1.2 * (.25 + .75 * len(t) / avg))
                 for w in tf)
        if sc > 0:
            res.append({"doc": name, "text": text.strip(), "score": round(sc, 2)})
    return sorted(res, key=lambda r: -r["score"])[:k]


# ---------------------------------------------------------------- cartella di lavoro
def safe_path(root, rel):
    """Percorso dentro la cartella di lavoro (symlink risolti). ValueError se ne esce o non c'è root."""
    if not root:
        raise ValueError("Nessuna cartella di lavoro autorizzata: sceglila in Impostazioni → Agente.")
    rootr = os.path.realpath(root)
    rel = str(rel or "").strip()
    if not rel or os.path.isabs(rel):
        raise ValueError("Usa un percorso relativo alla cartella di lavoro.")
    p = os.path.realpath(os.path.join(rootr, rel))
    if p != rootr and not p.startswith(rootr + os.sep):
        raise ValueError("Percorso fuori dalla cartella di lavoro: non consentito.")
    return p


TEXT_EXT = {".txt", ".md", ".markdown", ".csv", ".json", ".py", ".js", ".ts", ".html", ".css", ".xml", ".yaml",
            ".yml", ".sh", ".go", ".rs", ".java", ".c", ".h", ".cpp", ".log", ".tex", ".ini", ".toml", ".sql"}
MAX_READ = 40000  # caratteri restituiti al modello per un file


class Workspace:
    def __init__(self, root, extract):
        self.root = root
        self.extract = extract  # (nome, bytes) -> testo: riusa l'estrazione PDF/Word/Excel del backend

    def _rel(self, p):
        return os.path.relpath(p, os.path.realpath(self.root))

    def list_files(self, folder="."):
        base = safe_path(self.root, folder) if folder not in (".", "", "/") else os.path.realpath(self.root)
        if not os.path.isdir(base):
            raise ValueError("Cartella inesistente.")
        out = []
        for n in sorted(os.listdir(base))[:300]:
            if n.startswith("."): continue
            full = os.path.join(base, n)
            if os.path.isdir(full): out.append(self._rel(full) + "/")
            else: out.append(f"{self._rel(full)}  ({os.path.getsize(full)} B)")
        return "\n".join(out) or "(cartella vuota)"

    def read_file(self, path):
        p = safe_path(self.root, path)
        if not os.path.isfile(p):
            raise ValueError("File inesistente.")
        if os.path.getsize(p) > 80 * 1024 * 1024:
            raise ValueError("File troppo grande.")
        with open(p, "rb") as f:
            text = self.extract(os.path.basename(p), f.read())
        if len(text) > MAX_READ:
            return text[:MAX_READ] + f"\n\n[... troncato: mostrati {MAX_READ} caratteri su {len(text)}]"
        return text

    def grep(self, query, folder="."):
        base = safe_path(self.root, folder) if folder not in (".", "", "/") else os.path.realpath(self.root)
        q = (query or "").lower()
        if not q:
            raise ValueError("Testo da cercare mancante.")
        hits, seen = [], 0
        for dirpath, dirs, files in os.walk(base):
            dirs[:] = [d for d in dirs if not d.startswith(".") and d not in ("node_modules", "__pycache__")]
            for n in files:
                if os.path.splitext(n)[1].lower() not in TEXT_EXT: continue
                seen += 1
                if seen > 2000: break
                full = os.path.join(dirpath, n)
                try:
                    if os.path.getsize(full) > 2 * 1024 * 1024: continue
                    with open(full, encoding="utf-8", errors="ignore") as f:
                        for ln, line in enumerate(f, 1):
                            if q in line.lower():
                                hits.append(f"{self._rel(full)}:{ln}: {line.strip()[:200]}")
                                if len(hits) >= 60: return "\n".join(hits) + "\n[... altri risultati omessi]"
                except OSError:
                    continue
        return "\n".join(hits) or "Nessun risultato."

    def write_preview(self, path, content):
        p = safe_path(self.root, path)
        exists = os.path.exists(p)
        head = f"{self._rel(p)} · {len(content.encode())} B" + (" · SOVRASCRIVE il file esistente" if exists else " · nuovo file")
        return head + "\n\n" + content[:2000] + ("\n[...]" if len(content) > 2000 else "")

    def write_file(self, path, content):
        p = safe_path(self.root, path)
        os.makedirs(os.path.dirname(p), exist_ok=True)
        if p.lower().endswith(".docx"):
            _write_docx(p, content)
        else:
            tmp = p + ".gephid-tmp"
            with open(tmp, "w", encoding="utf-8") as f:
                f.write(content)
            os.replace(tmp, p)
        return f"File scritto: {self._rel(p)} ({os.path.getsize(p)} B)."


def _write_docx(path, md):
    """Markdown essenziale -> Word: titoli (#), elenchi (-, *, 1.), paragrafi, **grassetto**."""
    import docx
    d = docx.Document()
    for block in re.split(r"\n\s*\n", md.strip()):
        for line in block.splitlines():
            line = line.rstrip()
            if not line: continue
            h = re.match(r"(#{1,4})\s+(.*)", line)
            if h:
                d.add_heading(h.group(2).strip(), level=len(h.group(1))); continue
            b = re.match(r"\s*([-*•]|\d+[.)])\s+(.*)", line)
            par = d.add_paragraph(style=("List Number" if b and b.group(1)[0].isdigit() else "List Bullet") if b else None)
            text = b.group(2) if b else line
            for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
                if part: par.add_run(part).bold = bool(i % 2)
    d.save(path)


# ---------------------------------------------------------------- codice in sandbox
PROCS = {}       # id -> {"proc", "cmd", "dir", "started", "log"}
_PROCS_LOCK = threading.Lock()
RUN_TIMEOUT = 30


def _node_bin():
    cands = [shutil.which("node"), "/opt/homebrew/bin/node", "/usr/local/bin/node", os.path.expanduser("~/.volta/bin/node")]
    cands += sorted(glob.glob(os.path.expanduser("~/.nvm/versions/node/*/bin/node")), reverse=True)
    cands += sorted(glob.glob(os.path.expanduser("~/.local/share/fnm/node-versions/*/installation/bin/node")), reverse=True)
    cands += sorted(glob.glob(os.path.expanduser("~/Library/Application Support/fnm/node-versions/*/installation/bin/node")), reverse=True)
    return next((c for c in cands if c and os.path.isfile(c) and os.access(c, os.X_OK)), None)


def _profile(workdir):
    # rete: in uscita solo localhost; file: scrittura solo nella cartella temporanea del lavoro.
    w = os.path.realpath(workdir).replace('"', "")
    tmp = os.path.realpath(tempfile.gettempdir()).replace('"', "")
    return f"""(version 1)
(allow default)
(deny network-outbound)
(allow network-outbound (remote ip "localhost:*"))
(allow network-outbound (remote unix-socket))
(deny file-write*)
(allow file-write* (subpath "{w}") (subpath "{tmp}") (subpath "/dev"))
"""


def _listen_addrs(pid):
    """Indirizzi su cui il processo è in ascolto TCP, via lsof."""
    try:
        out = subprocess.run(["lsof", "-nP", "-a", "-iTCP", "-sTCP:LISTEN", "-p", str(pid), "-F", "n"],
                             capture_output=True, text=True, timeout=4).stdout
    except Exception:
        return []
    return sorted({ln[1:] for ln in out.splitlines() if ln.startswith("n")})  # es. "127.0.0.1:8097", "*:8098"


def _loopback_only(addrs):
    return all(a.startswith("127.") or a.startswith("[::1]") or a.startswith("localhost") for a in addrs)


def run_code(lang, code, timeout=RUN_TIMEOUT, background=False):
    """Esegue codice in sandbox. -> {"output","exit","timeout"} oppure, in background,
    {"id","running","output"}. Un processo in background che ascolta oltre 127.0.0.1 viene fermato."""
    lang = (lang or "python").lower()
    work = tempfile.mkdtemp(prefix="gephid-run-")
    if lang in ("python", "py", "python3"):
        import sys
        src = os.path.join(work, "main.py"); cmd = [sys.executable, "-u", src]
    elif lang in ("node", "javascript", "js"):
        node = _node_bin()
        if not node:
            return {"output": "Node.js non è installato su questo Mac: posso usare Python.", "exit": 127, "timeout": False}
        src = os.path.join(work, "main.js"); cmd = [node, src]
    else:
        return {"output": f"Linguaggio non supportato: {lang} (solo python o node).", "exit": 2, "timeout": False}
    with open(src, "w", encoding="utf-8") as f:
        f.write(code)
    prof = os.path.join(work, ".sandbox.sb")
    with open(prof, "w") as f:
        f.write(_profile(work))
    env = {"PATH": "/usr/bin:/bin", "HOME": work, "TMPDIR": work, "PYTHONDONTWRITEBYTECODE": "1", "LANG": "en_US.UTF-8"}
    full = ["/usr/bin/sandbox-exec", "-f", prof] + cmd
    if not background:
        try:
            r = subprocess.run(full, cwd=work, env=env, capture_output=True, text=True, timeout=timeout)
            out = (r.stdout + ("\n" + r.stderr if r.stderr else "")).strip()
            return {"output": out[-8000:] or "(nessun output)", "exit": r.returncode, "timeout": False}
        except subprocess.TimeoutExpired as e:
            out = ((e.stdout or b"").decode(errors="ignore") if isinstance(e.stdout, bytes) else (e.stdout or ""))
            return {"output": (out[-4000:] + f"\n[interrotto dopo {timeout}s]").strip(), "exit": -9, "timeout": True}
        finally:
            shutil.rmtree(work, ignore_errors=True)
    log = os.path.join(work, ".out.log")
    lf = open(log, "w")
    # stesso gruppo di processi del backend (niente nuova sessione): quando il launcher spegne il
    # backend col kill del gruppo, muoiono anche i processi lanciati dall'agente. sandbox-exec fa exec
    # del comando, quindi p è direttamente il processo Python/Node.
    p = subprocess.Popen(full, cwd=work, env=env, stdout=lf, stderr=subprocess.STDOUT)
    pid = uuid.uuid4().hex[:8]
    with _PROCS_LOCK:
        PROCS[pid] = {"proc": p, "cmd": lang, "dir": work, "started": time.time(), "log": log}
    time.sleep(1.5)  # tempo di avviarsi (e di mettersi in ascolto, se è un server)
    addrs = _listen_addrs(p.pid) if p.poll() is None else []
    if addrs and not _loopback_only(addrs):
        stop_proc(pid)
        return {"id": pid, "running": False, "output": "Fermato: era in ascolto su " + ", ".join(addrs) +
                ". Per sicurezza Gephid consente solo server su 127.0.0.1 (localhost)."}
    return {"id": pid, "running": p.poll() is None, "listen": addrs, "output": proc_output(pid)}


def proc_output(pid, tail=4000):
    e = PROCS.get(pid)
    if not e: return ""
    try:
        with open(e["log"], encoding="utf-8", errors="ignore") as f:
            return f.read()[-tail:]
    except OSError:
        return ""


def stop_proc(pid):
    with _PROCS_LOCK:
        e = PROCS.pop(pid, None)
    if not e: return False
    try:
        e["proc"].kill(); e["proc"].wait(timeout=3)
    except Exception: pass
    shutil.rmtree(e["dir"], ignore_errors=True)
    return True


def procs_status():
    out = []
    for pid, e in list(PROCS.items()):
        alive = e["proc"].poll() is None
        if alive:  # un server può mettersi in ascolto DOPO il controllo iniziale: ricontrolla
            addrs = _listen_addrs(e["proc"].pid)
            if addrs and not _loopback_only(addrs):
                stop_proc(pid); continue
        out.append({"id": pid, "lang": e["cmd"], "running": alive, "secs": int(time.time() - e["started"]),
                    "listen": _listen_addrs(e["proc"].pid) if alive else []})
    return out


def stop_all():
    for pid in list(PROCS):
        stop_proc(pid)


# ---------------------------------------------------------------- schemi per il template
def _fn(name, desc, props=None, required=None):
    return {"type": "function", "function": {"name": name, "description": desc, "parameters": {
        "type": "object", "properties": props or {}, "required": required or []}}}


S = {"type": "string"}
TOOLS_BASE = [
    _fn("calcola", "Calcola un'espressione aritmetica in modo esatto (+ - * / // % **, sqrt, ln, log10, sin, cos, pi). Usalo per OGNI conto.",
        {"espressione": S}, ["espressione"]),
    _fn("data_ora", "Restituisce data e ora attuali del computer."),
    _fn("cerca_nei_documenti", "Cerca nei documenti allegati alla chat i passaggi più pertinenti a una domanda (utile per documenti lunghi).",
        {"query": S}, ["query"]),
    _fn("esegui_codice", "Esegue un breve programma Python o Node in una sandbox senza rete. L'utente deve approvare. "
        "Usa in_background=true solo per processi che restano attivi (es. un server su 127.0.0.1).",
        {"linguaggio": {"type": "string", "enum": ["python", "node"]}, "codice": S, "in_background": {"type": "boolean"}},
        ["linguaggio", "codice"]),
]
TOOLS_FILES = [
    _fn("elenca_file", "Elenca file e sottocartelle nella cartella di lavoro dell'utente.", {"cartella": S}),
    _fn("leggi_file", "Legge un file della cartella di lavoro (testo, PDF, Word, Excel).", {"path": S}, ["path"]),
    _fn("cerca_nei_file", "Cerca un testo nei file di testo della cartella di lavoro.", {"testo": S, "cartella": S}, ["testo"]),
    _fn("scrivi_file", "Crea o sovrascrive un file nella cartella di lavoro (md, txt, csv, json, docx...). L'utente deve approvare.",
        {"path": S, "contenuto": S}, ["path", "contenuto"]),
]
CONFIRM = {"scrivi_file", "esegui_codice"}


def schemas(has_workspace):
    return TOOLS_BASE + (TOOLS_FILES if has_workspace else [])


def now_text():
    d = datetime.datetime.now()
    giorni = ["lunedì", "martedì", "mercoledì", "giovedì", "venerdì", "sabato", "domenica"]
    return f"{giorni[d.weekday()]} {d.strftime('%d/%m/%Y %H:%M')}"
